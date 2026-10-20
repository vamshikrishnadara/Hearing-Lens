"""Bounded community briefs from existing analysis; no files or result caches."""
from collections import Counter
from datetime import date
from io import BytesIO
import math
import unicodedata
from threading import Lock

_PDF_FONT_LOCK = Lock()


class BriefError(ValueError):
    """Safe export error without uploaded text."""


def _count(value):
    if type(value) is not int or not 0 <= value <= 5000:
        raise BriefError('Analysis counts are unavailable or inconsistent.')
    return value


def _number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _short(text, limit):
    text = ' '.join(''.join(c if not unicodedata.category(c).startswith('C') else ' '
                           for c in str(text)).split())
    if len(text) <= limit:
        return text
    prefix = text[:limit - 3]
    sentence = prefix.rfind('. ')
    boundary = prefix.rfind(' ')
    if sentence > len(prefix) // 2:
        prefix = prefix[:sentence]
    elif boundary > len(prefix) // 2:
        prefix = prefix[:boundary]
    return prefix.rstrip(' .') + '...'


def _safe(text, limit):
    # Subgroup labels can originate in uploaded cells, unlike redacted quotes.
    # Apply the same redactor to all selected prose before shortening it.
    from pipeline.redact import redact_text
    return _short(redact_text(str(text)).text, limit)


def _timeline_text(timeline):
    if not timeline or timeline.get('mode') != 'date':
        return 'Change over time is unavailable without comparable dates.'
    periods = []
    missing = 0
    for period in timeline.get('periods', []):
        if period['group'] == 'Missing period':
            missing += _count(period['total_comments'])
            continue
        try:
            day = date.fromisoformat(period['group'])
        except (ValueError, TypeError):
            raise BriefError('Timeline dates are invalid.') from None
        score = period.get('mean_sentiment')
        if _number(score) and -1 <= score <= 1 and period.get('scored_comments', 0) > 0:
            periods.append((day, score))
    periods.sort()
    if len(periods) < 2:
        text = 'At least two dated periods with scores are needed to compare tone.'
    else:
        first, last = periods[0], periods[-1]
        text = (f'Mean tone: {first[1]:+.2f} on {first[0]} to {last[1]:+.2f} on {last[0]} '
                '(first and last scored dates; scale -1 to +1).')
    if missing:
        text += f' {missing:,} undated comments excluded from this comparison.'
    return text


def _representation_text(gaps):
    fields = gaps.get('fields', [])
    if not fields:
        return 'No subgroup fields supplied; representation cannot be assessed.'
    valid = {f['field'] for f in fields if f.get('reference_status') == 'valid'
             and f.get('status') == 'comparison'}
    if not valid:
        return 'No usable reference comparison; the data cannot establish who was underrepresented.'
    candidates = [r for r in gaps.get('rows', [])
                  if r.get('field') in valid and not r.get('suppressed', True)
                  and r.get('flag') == 'underrepresented'
                  and _number(r.get('gap_percentage_points')) and r['gap_percentage_points'] < 0]
    candidates.sort(key=lambda r: (r['gap_percentage_points'], str(r['field']), str(r['group'])))
    if candidates:
        parts = [f'{_safe(r["group"], 24)} ({_safe(r["field"].removeprefix("subgroup__"), 18)}): '
                 f'{abs(r["gap_percentage_points"]):.1f} percentage points below reference'
                 for r in candidates[:2]]
        text = '; '.join(parts) + '.'
    else:
        text = 'No visible category met the underrepresentation threshold.'
    return text + ' Known category values only; hidden or unmatched groups are not assessed.'


def build_brief(result):
    """Select at most three themes and three question priorities without inference.

    Accept the pipeline result, not the uploaded table. Redaction resources may
    be reused; uploads, generated prose and file bytes are never globally cached.
    """
    try:
        size = _count(result['input_rows'])
        if not size:
            raise BriefError('Run analysis before preparing a brief.')
        core = result['themes']
        analyzed = _count(core['analyzed_comments'])
        ungrouped = _count(core['outlier_count'])
        themes = core['themes']
        counts = [_count(t['count']) for t in themes]
        ids = [t['theme_id'] for t in themes]
        if (analyzed > size or ungrouped > analyzed or sum(counts) + ungrouped != analyzed
                or len(set(ids)) != len(ids) or any(type(i) is not int or i <= 0 for i in ids)):
            raise BriefError('Theme counts are inconsistent.')
        selected = []
        for theme in sorted(themes, key=lambda t: (-t['count'], t['theme_id']))[:3]:
            quote = next(iter(theme.get('quotes', [])), None)
            selected.append({'label': _safe(theme['label'], 56), 'count': theme['count'],
                             'share': theme['count'] / analyzed if analyzed else 0,
                             'quote': _safe(quote['text'], 140) if quote else None})
        rows = result['affect']['rows']
        if len(rows) != size or {r['row_position'] for r in rows} != set(range(size)):
            raise BriefError('Sentiment rows are inconsistent.')
        scored = [r for r in rows if r['status'] == 'scored']
        labels = Counter(r['sentiment'] for r in scored)
        if set(labels) - {'positive', 'neutral', 'negative'}:
            raise BriefError('Sentiment labels are invalid.')
        if scored:
            sentiment = ', '.join(f'{labels[label] / len(scored):.0%} {label}'
                                  for label in ('negative', 'neutral', 'positive'))
            sentiment += f' among {len(scored):,} scored comments; {size - len(scored):,} unscored.'
        else:
            sentiment = 'No sentiment scores available; unscored comments are not neutral.'
        questions = result['questions']
        groups = {g['question_group_id']: g for g in questions['groups']}
        priorities = questions['top_unanswered']
        if len(set(priorities)) != len(priorities) or any(i not in groups for i in priorities):
            raise BriefError('Question priorities are inconsistent.')
        chosen = []
        for qid in priorities[:3]:
            group = groups[qid]
            chosen.append({'text': _safe(group['representative']['text'], 150),
                           'count': _count(group['comment_count'])})
        if not groups:
            question_note = 'No question candidates detected; concerns may still be present.'
        elif questions.get('response_status') != 'available':
            question_note = 'No usable response supplied; these are not confirmed unanswered questions.'
        elif not chosen:
            question_note = 'All groups have possible response matches; this does not confirm answers.'
        else:
            question_note = 'Missing or partial response matches; inspect passages before claiming no answer.'
        return {'title': 'Community feedback brief',
                'summary': f'{size:,} cleaned comments analyzed to support discussion. '
                           'Counts describe comments, not necessarily unique people.',
                'theme_note': f'Top {len(selected)} of {len(themes)} themes by comment count; '
                              f'{ungrouped:,} ungrouped, {size - analyzed:,} excluded. '
                              'Shares use theme-analyzed comments.',
                'themes': selected, 'sentiment': sentiment,
                'timeline': _timeline_text(result.get('timeline')),
                'representation': _representation_text(result.get('gaps', {})),
                'questions': chosen, 'question_note': question_note,
                'method': 'Method: embedding-based themes, English sentiment models and question rules. '
                          'Up to three themes, one quote excerpt each, and three questions shown; "..." marks shortened text. '
                          'Automated findings need review. Redaction can miss identifiers.'}
    except BriefError:
        raise
    except (KeyError, TypeError, ValueError, AttributeError):
        raise BriefError('The analysis is incomplete; no brief was generated.') from None


def _comment_count(count):
    return f'{count:,} comment' + ('' if count == 1 else 's')


def brief_paragraphs(brief):
    """Shared prose for Word and PDF; each row is (style, prefix, text)."""
    rows = [('Title', '', brief['title']), ('Normal', '', brief['summary']),
            ('Heading 1', '', 'Leading themes and quote excerpts'),
            ('Normal', '', brief['theme_note'])]
    for theme in brief['themes']:
        prefix = f'{theme["label"]} - {_comment_count(theme["count"])} ({theme["share"]:.1%}). '
        text = f'Excerpt: "{theme["quote"]}"' if theme['quote'] else 'No quote available.'
        rows.append(('Normal', prefix, text))
    if not brief['themes']:
        rows.append(('Normal', '', 'No themes were identified.'))
    rows.extend([('Heading 1', '', 'Sentiment and change over time'),
                 ('Normal', '', brief['sentiment'] + ' ' + brief['timeline']),
                 ('Heading 1', '', 'Representation'), ('Normal', '', brief['representation']),
                 ('Heading 1', '', 'Priority questions for follow up'),
                 ('Normal', '', brief['question_note'])])
    rows.extend(('Normal', '', f'{i}. {q["text"]} ({_comment_count(q["count"])})')
                for i, q in enumerate(brief['questions'], 1))
    rows.append(('Method', '', brief['method']))
    return rows


def docx_bytes(brief):
    """Render the bounded brief as an editable, in-memory Word document."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.55)
    section.left_margin = section.right_margin = Inches(.65)
    for name in ('Normal', 'Title', 'Heading 1'):
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(11 if name != 'Title' else 16)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.space_before = Pt(0 if name != 'Heading 1' else 5)
    document.styles['Heading 1'].font.bold = True
    for style, prefix, text in brief_paragraphs(brief):
        p = document.add_paragraph(style='Normal' if style == 'Method' else style)
        if prefix:
            p.add_run(prefix).bold = True
        p.add_run(text)
        if style == 'Method':
            p.paragraph_format.space_before = Pt(5)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.title = brief['title']
    document.core_properties.subject = 'Community feedback analysis'
    document.core_properties.comments = ''
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(brief):
    """Render one page in memory, with embedded fonts and no external renderer."""
    from pathlib import Path
    from xml.sax.saxutils import escape
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    # These redistributable font files ship with ReportLab on each platform.
    # Register font resources only; never cache a brief or uploaded content.
    font_root = Path(reportlab.__file__).parent / 'fonts'
    with _PDF_FONT_LOCK:
        for name, filename in [('HLVera', 'Vera.ttf'), ('HLVeraBold', 'VeraBd.ttf')]:
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(font_root / filename)))
        pdfmetrics.registerFontFamily('HLVera', normal='HLVera', bold='HLVeraBold')
    paragraphs = brief_paragraphs(brief)
    for style, prefix, text in paragraphs:
        for font, value in [('HLVeraBold', prefix), ('HLVeraBold' if style == 'Heading 1' else 'HLVera', text)]:
            glyphs = pdfmetrics.getFont(font).face.charToGlyph
            if any(ord(c) not in glyphs for c in value if not c.isspace()):
                raise BriefError('The PDF font cannot display some characters. Download the Word brief instead.')
    styles = {name: ParagraphStyle(name, fontName='HLVeraBold' if name == 'Heading 1' else 'HLVera',
        fontSize=16 if name == 'Title' else 11, leading=19 if name == 'Title' else 13,
        spaceAfter=3, spaceBefore=5 if name in ('Heading 1', 'Method') else 0,
        keepWithNext=name in ('Title', 'Heading 1')) for name in ('Title', 'Normal', 'Heading 1', 'Method')}
    story = [Paragraph(('<b>' + escape(prefix) + '</b>' if prefix else '') + escape(text), styles[style])
             for style, prefix, text in paragraphs]
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=(612, 792), topMargin=39.6, bottomMargin=39.6,
                           leftMargin=46.8, rightMargin=46.8, title=brief['title'], author='',
                           subject='Community feedback analysis', creator='Hearing Lens')
    doc.build(story)
    if doc.page != 1:
        raise BriefError('The PDF brief could not fit on one page. Download the Word brief instead.')
    return output.getvalue()


def community_docx(result):
    return docx_bytes(build_brief(result))


def community_pdf(result):
    return pdf_bytes(build_brief(result))
